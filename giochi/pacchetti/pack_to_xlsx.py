"""Scrive un pacchetto contenuti SGS Learning (JSON) nel modello Excel di revisione."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

SLOTS = [('A1', 'Atrio · totem 1'), ('A2', 'Atrio · totem 2'), ('A3', 'Atrio · totem 3'), ('A4', 'Atrio · totem 4 (verso le scale)'), ('A5', 'Atrio · totem 5 (verso le scale)'),
         ('W1', 'Aula · parete sinistra'), ('W2', 'Aula · parete destra'), ('P1', 'Banchina · totem 1'), ('P2', 'Banchina · totem 2'), ('P3', 'Banchina · totem 3'),
         ('P4', 'Banchina · totem 4'), ('P5', 'Banchina · totem 5, al cancello del cantiere')]
DIVISE = [('polo blu', 'Polo a maniche corte'), ('camicia bianca', 'Camicia a maniche lunghe'), ('completo scuro', 'Giacca e pantaloni scuri'), ('giacca grigia', 'Giacca grigia'),
          ('maglia verde petrolio', 'Maglia a maniche lunghe'), ('maglia bordeaux', 'Maglia a maniche lunghe'), ('divisa condotta', 'Divisa del personale di condotta'),
          ('divisa stazione', 'Divisa del personale di stazione'), ('lavoro arancio', 'Abbigliamento da officina con bande'), ('alta visibilità', 'Giubbotto ad alta visibilità')]
CAPELLI = ['corti', 'rasati', 'ricci', 'lunghi', 'caschetto', 'chignon', 'coda', 'calvo']
COL_CAPELLI = ['nero', 'castano', 'biondo', 'grigio', 'rosso']
COLORI = ['verde acqua', 'verde', 'blu', 'azzurro', 'arancio', 'ambra', 'viola', 'grigio', 'rosso']
POSE = ['in piedi', 'legge', 'telefono']
MODELLI = [('sequenza', '2–6 riquadri collegati da frecce (fasi, evoluzione nel tempo)'), ('colonne', '2–4 colonne affiancate con titolo e testo'), ('ciclo', '3–6 fasi in cerchio; la «Nota sul pannello» va al centro'),
           ('confronto', 'Tabella a due colonne: la prima riga contiene le intestazioni in Testo e Testo 2, le altre etichetta, sinistra, destra'), ('elenco', '2–6 riquadri in griglia'),
           ('anelli', '3–6 cerchi concentrici dall\'esterno al centro'), ('scheda', 'Scheda di ruolo: 1ª riga riquadro (es. Mission), righe intermedie voci con spunta, ultima riga fascia colorata')]
DEDICATE = [f'M1-{i:02d}' for i in range(1, 13)]
TIPI_LAB = [('ordina', 'Mettere le voci in sequenza. Risposta = posizione corretta (1, 2, 3…)'), ('abbina', 'Associare ogni voce a un\'opzione. Opzioni nel foglio Laboratori separate da «|»; Risposta = testo esatto dell\'opzione'),
            ('scelta', 'Scelta multipla: selezionare tutte e sole le voci giuste. Risposta = Sì oppure No; Spiegazione mostrata per le voci errate')]
OGGETTI = [('rotaie', 'Cantiere binario 1 · rotaie di scorta', 'sì'), ('segnale', 'Imbocco galleria · segnale luminoso', 'sì'), ('sse', 'Galleria · porta della sottostazione elettrica', 'sì'),
           ('fessurimetro', 'Galleria · fessurimetro sul piedritto', 'sì'), ('estintore', 'Banchina · estintore sul pilastro', 'no'), ('dae', 'Atrio · defibrillatore DAE', 'no'), ('sos', 'Banchina · colonnina SOS', 'no'),
           ('tornelli', 'Atrio · tornelli', 'no'), ('lineagialla', 'Banchina · linea gialla e bordo banchina', 'no'), ('display', 'Banchina · display «prossimo treno»', 'no')]
REVISIONE = ['Da verificare', 'Approvato', 'Da correggere']

NAVY, SKY, REVFILL, EXFILL = '15386E', 'DCEFFA', 'EEF1F6', 'FFF7E0'
F = lambda **k: Font(name='Arial', **k)
thin = Side(style='thin', color='C9D3DD')

def sheet(wb, name, headers, widths, rows, review=False, notes=None, tab=None):
    ws = wb.create_sheet(name)
    if tab: ws.sheet_properties.tabColor = tab
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h); cell.font = F(bold=True, color='FFFFFF', size=10); cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = Border(bottom=thin)
        if notes and h in notes: cell.comment = Comment(notes[h], 'SGS Learning')
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.font = F(size=10); cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.border = Border(bottom=thin)
            if review and headers[c - 1] in ('Revisione', 'Note revisore'): cell.fill = PatternFill('solid', fgColor=REVFILL)
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30
    return ws

def dv_list(ws, formula, col, n=400, allow_blank=True, prompt=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=allow_blank, showErrorMessage=True, errorTitle='Valore non previsto', error='Scegli un valore dall\'elenco.')
    if prompt: dv.promptTitle = 'Valori ammessi'; dv.prompt = prompt; dv.showInputMessage = True
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}{n}')

def build(pack, out, example=False):
    wb = Workbook(); wb.remove(wb.active); M = pack['modulo']
    # ---- Istruzioni
    ws = wb.create_sheet('Istruzioni'); ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 70; ws.column_dimensions['C'].width = 18
    lines = [('SGS Learning · pacchetto contenuti', None), (f'{M["titolo"]} · {M["nome"]}', None), ('', None),
      ('Come si usa', 'h'),
      ('1. Compila o correggi', 'I fogli Modulo, Tappe, Punti, Pannelli, Laboratori, Voci, Quiz, Sopralluogo, Esame, Referenti e Sigle. Le intestazioni blu non vanno modificate; l\'ordine delle colonne sì, se serve.'),
      ('2. Revisiona', 'Nelle colonne grigie «Revisione» scegli Approvato, Da correggere o Da verificare e scrivi le osservazioni in «Note revisore». Il gioco ignora queste colonne.'),
      ('3. Approva', 'Nel foglio Modulo porta «Stato» su Approvato e indica chi approva. Finché resta Bozza il gioco mostra la dicitura BOZZA.'),
      ('4. Carica', 'Nel gioco: schermata iniziale › «Carica modulo» › scegli questo file .xlsx. Se ci sono errori il gioco li elenca con foglio e riga.'),
      ('', None), ('Regole principali', 'h'),
      ('Tappe', 'Da 1 a 12. Ogni tappa occupa una posizione diversa della stazione (foglio Elenchi, colonna Posizione). Numero tappa: 1, 2, 3… senza salti.'),
      ('Quiz', 'Almeno 1 per tappa, consigliati 2. Da 2 a 4 risposte: una corretta e fino a tre errate.'),
      ('Esame', 'Almeno tante domande quante indicate in «Domande esame»; meglio 2–3 per tappa, così ogni prova è diversa. Quattro risposte per domanda.'),
      ('Laboratori', 'Da 0 a 3 per tappa. Voci: «ordina» 3–6, «abbina» 2–8 con 2–6 opzioni, «scelta» 3–8.'),
      ('Sopralluogo', 'Per le tappe di tipo «sopralluogo»: 1–6 oggetti del catalogo. Gli oggetti in cantiere o galleria richiedono gilet e casco (armadietto in aula) e l\'apertura del cancello da parte del referente.'),
      ('Grafica del pannello', '«dedicata» usa le infografiche disegnate apposta per il Modulo 1 (codici M1-01…M1-12; i testi del foglio Pannelli vengono ignorati). «modello» costruisce il pannello dal foglio Pannelli con il modello scelto.'),
      ('Referenti', 'Indicati solo per ruolo. Nessun nominativo reale. Il referente guida accoglie l\'allievo; l\'esaminatore sta in aula e conduce l\'esame.'),
      ('Testi', 'Nei testi dei dialoghi si possono evidenziare parole con <em>parola</em>.'),
      ('', None), ('Valori ammessi', 'h')]
    r = 1
    for a, b in lines:
        if b == 'h': ws.cell(r, 1, a).font = F(bold=True, size=12, color=NAVY)
        elif b is None: ws.cell(r, 1, a).font = F(bold=(r <= 2), size=(16 if r == 1 else 11), color=NAVY)
        else:
            ws.cell(r, 1, a).font = F(bold=True, size=10); c = ws.cell(r, 2, b); c.font = F(size=10); c.alignment = Alignment(wrap_text=True, vertical='top'); ws.cell(r, 1).alignment = Alignment(vertical='top')
        r += 1
    def table(title, rows, head):
        nonlocal r
        ws.cell(r, 1, title).font = F(bold=True, size=10, color=NAVY); r += 1
        for c, h in enumerate(head, 1): cell = ws.cell(r, c, h); cell.font = F(bold=True, size=9, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='44505E')
        r += 1
        for row in rows:
            for c, v in enumerate(row, 1): cell = ws.cell(r, c, v); cell.font = F(size=9); cell.alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
        r += 1
    table('Posizione della tappa', SLOTS, ['Codice', 'Dove si trova'])
    table('Modello del pannello', MODELLI, ['Modello', 'Uso'])
    table('Tipo di laboratorio', TIPI_LAB, ['Tipo', 'Come si compila'])
    table('Oggetti per il sopralluogo', [(a, b, c) for a, b, c in OGGETTI], ['Oggetto', 'Dove', 'Serve DPI'])
    table('Divise dei referenti', DIVISE, ['Divisa', 'Descrizione'])
    ws.cell(r, 1, 'Capelli: ' + ', '.join(CAPELLI) + ' · Colore capelli: ' + ', '.join(COL_CAPELLI) + ' · Colore distintivo: ' + ', '.join(COLORI) + ' · Posa: ' + ', '.join(POSE) + ' · Carnagione: da 1 (chiara) a 6 (scura)').font = F(size=9)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); ws.row_dimensions[r].height = 42
    ws.sheet_view.showGridLines = False

    # ---- Modulo
    rows = [('ID modulo', M['id'], 'Codice breve, senza spazi (es. M1). Serve a tenere separati i risultati dei diversi moduli.'),
            ('Titolo', M['titolo'], 'Es. Modulo 1'), ('Nome', M['nome'], 'Titolo completo del modulo'), ('Nome breve', M['breve'], 'Sottotitolo nella schermata iniziale'),
            ('Corso', M['corso'], ''), ('Ruolo allievo', M['allievo'], 'Come appare il protagonista nei dialoghi')]
    for k in range(4):
        s = M['sottomoduli'][k] if k < len(M['sottomoduli']) else {'codice':'', 'titolo':''}
        rows.append((f'Sotto-modulo {k + 1} · codice', s['codice'], 'Es. 1.1 (lasciare vuoto se non serve)')); rows.append((f'Sotto-modulo {k + 1} · titolo', s['titolo'], ''))
    rows += [('Revisione', M['revisione'], ''), ('Data', M['data'], 'AAAA-MM-GG'), ('Stato', M['stato'], 'Bozza oppure Approvato'), ('Approvato da', M['approvatoDa'], 'Ruolo di chi approva, es. RSGS'),
             ('Fonti', M['fonti'], 'Documenti da cui sono tratti i contenuti'), ('Referente guida', M['guida'], 'ID dal foglio Referenti'), ('Esaminatore', M['esaminatore'], 'ID dal foglio Referenti'),
             ('Domande esame', M['domande'], 'Numero di domande della prova finale'), ('Soglia esame', M['soglia'], 'Risposte corrette necessarie'),
             ('Presentazione', M['presentazione'], 'Testo della schermata iniziale'), ('Apertura', M['apertura'], 'Prima battuta del protagonista all\'inizio')]
    for k in range(3): rows.append((f'Spiegazione guida {k + 1}', M['guidaSpiegazione'][k] if k < len(M['guidaSpiegazione']) else '', 'Cosa dice il referente guida quando l\'allievo si presenta'))
    wsM = sheet(wb, 'Modulo', ['Campo', 'Valore', 'Note'], [26, 80, 46], rows)
    r0 = len(rows) + 3
    wsM.cell(r0, 1, 'Controllo revisione').font = F(bold=True, size=10, color=NAVY)
    checks = [('Righe «Da correggere»', '=COUNTIF(Tappe!Q:Q,"Da correggere")+COUNTIF(Punti!D:D,"Da correggere")+COUNTIF(Laboratori!G:G,"Da correggere")+COUNTIF(Quiz!J:J,"Da correggere")+COUNTIF(Sopralluogo!H:H,"Da correggere")+COUNTIF(Esame!J:J,"Da correggere")'),
              ('Righe «Da verificare»', '=COUNTIF(Tappe!Q:Q,"Da verificare")+COUNTIF(Punti!D:D,"Da verificare")+COUNTIF(Laboratori!G:G,"Da verificare")+COUNTIF(Quiz!J:J,"Da verificare")+COUNTIF(Sopralluogo!H:H,"Da verificare")+COUNTIF(Esame!J:J,"Da verificare")'),
              ('Righe «Approvato»', '=COUNTIF(Tappe!Q:Q,"Approvato")+COUNTIF(Punti!D:D,"Approvato")+COUNTIF(Laboratori!G:G,"Approvato")+COUNTIF(Quiz!J:J,"Approvato")+COUNTIF(Sopralluogo!H:H,"Approvato")+COUNTIF(Esame!J:J,"Approvato")'),
              ('Tappe', '=COUNT(Tappe!A:A)'), ('Domande d\'esame disponibili', '=COUNTA(Esame!C:C)-1')]
    for i, (a, f) in enumerate(checks):
        wsM.cell(r0 + 1 + i, 1, a).font = F(size=10); c = wsM.cell(r0 + 1 + i, 2, f); c.font = F(size=10, bold=True); c.alignment = Alignment(horizontal='left')
    dv = DataValidation(type='list', formula1='"Bozza,Approvato"', allow_blank=False); wsM.add_data_validation(dv)
    dv.add(f'B{[i for i, x in enumerate(rows, 2) if x[0] == "Stato"][0]}')

    T = pack['tappe']
    rev = 'Da verificare'
    # ---- Tappe
    rows = []
    for t in T:
        sp = t['spiegazione'] + [''] * (3 - len(t['spiegazione']))
        rows.append([t['n'], t['id'], t['sotto'], t['titolo'], t['slot'], t['referente'], t['tipo'], t['grafica'], t.get('dedicata', ''), t['modello'], t['notaPannello'], sp[0], sp[1], sp[2], t['nota'], None, rev, ''])
    wsT = sheet(wb, 'Tappe', ['N', 'ID', 'Sotto-modulo', 'Titolo', 'Posizione', 'Referente', 'Tipo', 'Grafica', 'Grafica dedicata', 'Modello pannello', 'Nota sul pannello',
                              'Spiegazione 1', 'Spiegazione 2', 'Spiegazione 3', 'Nota critica', 'Controlli', 'Revisione', 'Note revisore'],
                [5, 13, 11, 30, 10, 11, 12, 11, 11, 13, 40, 50, 50, 30, 50, 22, 14, 30], rows, review=True,
                notes={'Posizione':'Codice della posizione in stazione (vedi Istruzioni). Ogni tappa in una posizione diversa.', 'Referente':'ID dal foglio Referenti',
                       'Spiegazione 1':'Cosa dice il referente prima del laboratorio. Fino a tre battute.', 'Nota critica':'Facoltativa: punti delle fonti da precisare.', 'Controlli':'Calcolato: punti, quiz, domande d\'esame per tappa'})
    for i in range(len(T)):
        rr = i + 2; wsT.cell(rr, 16, f'=COUNTIF(Punti!A:A,A{rr})&" punti · "&COUNTIF(Quiz!A:A,A{rr})&" quiz · "&COUNTIF(Esame!B:B,A{rr})&" esame"').font = F(size=9, color='44505E')
    dv_list(wsT, 'Elenchi!$A$2:$A$13', 'E'); dv_list(wsT, 'Referenti!$A$2:$A$40', 'F'); dv_list(wsT, '"standard,sopralluogo"', 'G'); dv_list(wsT, '"dedicata,modello"', 'H')
    dv_list(wsT, 'Elenchi!$B$2:$B$13', 'I'); dv_list(wsT, 'Elenchi!$C$2:$C$8', 'J'); dv_list(wsT, 'Elenchi!$I$2:$I$4', 'Q')
    # ---- Punti
    rows = [[t['n'], k + 1, p, rev, ''] for t in T for k, p in enumerate(t['punti'])]
    wsP = sheet(wb, 'Punti', ['Tappa', 'Ordine', 'Punto chiave', 'Revisione', 'Note revisore'], [7, 7, 100, 14, 30], rows, review=True); dv_list(wsP, 'Elenchi!$I$2:$I$4', 'D', 800)
    # ---- Pannelli
    rows = [[t['n'], k + 1, e['titolo'], e['testo'], e['testo2']] for t in T for k, e in enumerate(t['elementi'])]
    sheet(wb, 'Pannelli', ['Tappa', 'Ordine', 'Titolo', 'Testo', 'Testo 2'], [7, 7, 34, 60, 50], rows, notes={'Testo 2':'Solo per il modello «confronto» (colonna destra).'})
    # ---- Laboratori e voci
    rows, vrows = [], []
    for t in T:
        for k, L in enumerate(t['labs']):
            rows.append([t['n'], k + 1, L['tipo'], L['titolo'], L['istruzioni'], ' | '.join(L['opzioni']), rev, ''])
            for v in L['voci']: vrows.append([t['n'], k + 1, v['testo'], v['risposta'], v['spiegazione']])
    wsL = sheet(wb, 'Laboratori', ['Tappa', 'Lab', 'Tipo', 'Titolo', 'Istruzioni', 'Opzioni', 'Revisione', 'Note revisore'], [7, 6, 10, 34, 50, 60, 14, 30], rows, review=True,
                notes={'Opzioni':'Solo per «abbina»: opzioni separate da «|».'}); dv_list(wsL, '"ordina,abbina,scelta"', 'C'); dv_list(wsL, 'Elenchi!$I$2:$I$4', 'G')
    sheet(wb, 'Voci', ['Tappa', 'Lab', 'Voce', 'Risposta', 'Spiegazione'], [7, 6, 70, 34, 50], vrows,
          notes={'Risposta':'ordina: posizione (1, 2, 3…) · abbina: testo esatto dell\'opzione · scelta: Sì oppure No', 'Spiegazione':'Solo per «scelta»: mostrata se l\'allievo seleziona una voce errata.'})
    # ---- Quiz
    rows = []
    for t in T:
        for k, q in enumerate(t['quiz']):
            e = q['errate'] + [''] * (3 - len(q['errate'])); rows.append([t['n'], k + 1, q['domanda'], q['corretta'], e[0], e[1], e[2], q['spiegazione'], q['fonte'], rev, ''])
    wsQ = sheet(wb, 'Quiz', ['Tappa', 'Ordine', 'Domanda', 'Risposta corretta', 'Errata 1', 'Errata 2', 'Errata 3', 'Spiegazione', 'Fonte', 'Revisione', 'Note revisore'],
                [7, 7, 44, 40, 34, 34, 22, 54, 12, 14, 30], rows, review=True); dv_list(wsQ, 'Elenchi!$I$2:$I$4', 'J')
    # ---- Sopralluogo
    rows = [[t['n'], h['oggetto'], h['etichetta'], h['titolo'], h['descrizione'], ' | '.join(h['opzioni']), h['corretta'], rev, ''] for t in T for h in t['sopralluogo']]
    wsS = sheet(wb, 'Sopralluogo', ['Tappa', 'Oggetto', 'Azione', 'Titolo', 'Descrizione', 'Opzioni', 'Risposta corretta', 'Revisione', 'Note revisore'], [7, 13, 32, 30, 60, 50, 24, 14, 30], rows, review=True,
                notes={'Oggetto':'Dal catalogo (vedi Istruzioni).', 'Azione':'Testo del suggerimento a schermo, es. «Esamina il segnale».', 'Opzioni':'Separate da «|».'})
    dv_list(wsS, 'Elenchi!$D$2:$D$11', 'B'); dv_list(wsS, 'Elenchi!$I$2:$I$4', 'H')
    # ---- Esame
    rows = [[f'E{k + 1:02d}', e['tappa'], e['domanda'], e['corretta'], *(e['errate'] + [''] * (3 - len(e['errate']))), e['spiegazione'], e['fonte'], rev, ''] for k, e in enumerate(pack['esame'])]
    wsE = sheet(wb, 'Esame', ['ID', 'Tappa', 'Domanda', 'Risposta corretta', 'Errata 1', 'Errata 2', 'Errata 3', 'Spiegazione', 'Fonte', 'Revisione', 'Note revisore'],
                [6, 7, 44, 40, 32, 32, 32, 50, 12, 14, 30], rows, review=True); dv_list(wsE, 'Elenchi!$I$2:$I$4', 'J')
    # ---- Referenti
    rows = [[x['id'], x['ruolo'], x['struttura'], x['sigla'], x['colore'], x['divisa'], x['figura'], x['capelli'], x['coloreCapelli'], x['carnagione'], x['posa'], x['casco'], x['saluto'], x['domandaExtra'], x['rispostaExtra']] for x in pack['referenti']]
    wsR = sheet(wb, 'Referenti', ['ID', 'Ruolo', 'Struttura', 'Sigla', 'Colore', 'Divisa', 'Figura', 'Capelli', 'Colore capelli', 'Carnagione', 'Posa', 'Casco', 'Saluto', 'Domanda extra', 'Risposta extra'],
                [9, 24, 30, 7, 12, 20, 7, 11, 12, 10, 10, 7, 50, 34, 50], rows, notes={'Ruolo':'Solo ruoli, nessun nominativo.', 'Sigla':'Massimo 4 caratteri, compare nel distintivo del dialogo.', 'Domanda extra':'Facoltativa: una domanda in più nel dialogo, con la sua risposta.'})
    dv_list(wsR, 'Elenchi!$G$2:$G$10', 'E', 40); dv_list(wsR, 'Elenchi!$E$2:$E$11', 'F', 40); dv_list(wsR, '"F,M"', 'G', 40); dv_list(wsR, 'Elenchi!$F$2:$F$9', 'H', 40)
    dv_list(wsR, 'Elenchi!$H$2:$H$6', 'I', 40); dv_list(wsR, '"in piedi,legge,telefono"', 'K', 40); dv_list(wsR, '"sì,no"', 'L', 40)
    dvn = DataValidation(type='whole', operator='between', formula1='1', formula2='6', showErrorMessage=True, error='Da 1 a 6'); wsR.add_data_validation(dvn); dvn.add('J2:J40')
    # ---- Sigle
    sheet(wb, 'Sigle', ['Sigla', 'Significato'], [16, 90], [[g['sigla'], g['significato']] for g in pack['sigle']])
    # ---- Elenchi (nascosto, alimenta i menu a tendina)
    wl = wb.create_sheet('Elenchi')
    cols = [('Posizione', [s[0] for s in SLOTS]), ('Grafica dedicata', DEDICATE), ('Modello', [m[0] for m in MODELLI]), ('Oggetto', [o[0] for o in OGGETTI]), ('Divisa', [d[0] for d in DIVISE]),
            ('Capelli', CAPELLI), ('Colore', COLORI), ('Colore capelli', COL_CAPELLI), ('Revisione', REVISIONE)]
    for c, (h, vals) in enumerate(cols, 1):
        wl.cell(1, c, h).font = F(bold=True, size=10)
        for i, v in enumerate(vals, 2): wl.cell(i, c, v).font = F(size=10)
    wl.sheet_state = 'hidden'
    wb.properties.title = f'SGS Learning · {M["titolo"]}'; wb.properties.creator = 'SGS Learning'
    wb.save(out)

if __name__ == '__main__':
    pack = json.load(open(sys.argv[1], encoding='utf-8')); build(pack, sys.argv[2]); print('ok', sys.argv[2])
